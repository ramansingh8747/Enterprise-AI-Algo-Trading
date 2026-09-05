import json
import re
from pathlib import Path
from typing import Any


class StrategyFileParser:
    """Extracts explicit strategy text/rules without inventing trading logic."""

    MAX_TEXT_CHARS = 120_000
    KEY_ALIASES = {
        "symbol": "symbol",
        "instrument": "symbol",
        "timeframe": "timeframe",
        "interval": "timeframe",
        "entry": "entry",
        "entry condition": "entry",
        "entry conditions": "entry",
        "exit": "exit",
        "exit condition": "exit",
        "exit conditions": "exit",
        "stop loss": "stop_loss",
        "stoploss": "stop_loss",
        "target": "target",
        "take profit": "target",
        "quantity": "quantity",
        "order type": "order_type",
        "side": "side",
        "indicator": "indicator",
    }

    def parse(self, path: Path, extension: str) -> tuple[str, dict[str, Any], list[str]]:
        extension = extension.lower()
        if extension == ".xlsx":
            text, structured = self._parse_xlsx(path)
        elif extension == ".docx":
            text = self._parse_docx(path)
            structured = {}
        elif extension == ".pdf":
            text = self._parse_pdf(path)
            structured = {}
        elif extension == ".txt":
            text = path.read_text(encoding="utf-8", errors="replace")
            structured = {}
        else:
            raise ValueError("Unsupported strategy file type.")

        text = text[: self.MAX_TEXT_CHARS]
        extracted = dict(structured)
        extracted.update(self._extract_key_value_rules(text))
        if not extracted:
            warnings = ["No explicit key/value strategy rules were detected. Review the extracted text before confirmation."]
        else:
            warnings = []
        if len(text) >= self.MAX_TEXT_CHARS:
            warnings.append(f"Extracted text was truncated to {self.MAX_TEXT_CHARS} characters.")
        return text, extracted, warnings

    def _parse_xlsx(self, path: Path) -> tuple[str, dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel import requires openpyxl to be installed.") from exc

        workbook = load_workbook(path, read_only=True, data_only=True)
        lines: list[str] = []
        structured: dict[str, Any] = {}
        for sheet in workbook.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if not values:
                    continue
                lines.append(" | ".join(values))
                if len(values) >= 2:
                    key = self._normalise_key(values[0])
                    if key:
                        structured[key] = values[1] if len(values) == 2 else values[1:]
        workbook.close()
        return "\n".join(lines), structured

    @staticmethod
    def _parse_docx(path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX import requires python-docx to be installed.") from exc
        document = Document(path)
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    paragraphs.append(" | ".join(values))
        return "\n".join(paragraphs)

    @staticmethod
    def _parse_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF import requires pypdf to be installed.") from exc
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    def _extract_key_value_rules(self, text: str) -> dict[str, Any]:
        extracted: dict[str, Any] = {}
        for raw_line in text.splitlines():
            line = raw_line.strip().strip("|-")
            if not line:
                continue
            parts = re.split(r"\s*(?:\||:|=|\t)\s*", line, maxsplit=1)
            if len(parts) != 2:
                continue
            key = self._normalise_key(parts[0])
            value = parts[1].strip()
            if key and value:
                extracted[key] = value
        return extracted

    def _normalise_key(self, value: str) -> str | None:
        key = re.sub(r"\s+", " ", value.strip().lower())
        return self.KEY_ALIASES.get(key)
